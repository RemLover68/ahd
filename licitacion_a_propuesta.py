#!/usr/bin/env python3
"""
licitacion_a_propuesta.py

Pipeline local-first para procesar una licitación con RAG especializado:

  Fase 0  Convertir documentos a texto y construir un corpus local.
  Fase 1  Extraer la lista consolidada de ítems de la licitación.
  Fase 2  Proponer el índice de secciones/subsistemas de la propuesta.
  Fase 3  Generar, por sección, requisitos y propuesta con citas locales.
  Fase 4  Generar resumen ejecutivo y guardar result.json + .docx.

Este módulo habla con un servidor local OpenAI-compatible (Ollama,
llama.cpp server, vLLM, LM Studio, etc.) y usa embeddings locales para RAG.
"""

from __future__ import annotations

import argparse
import asyncio
import datetime as _dt
import json
import os
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
from openai import AsyncOpenAI
from sentence_transformers import SentenceTransformer

# --------------------------------------------------------------------------- #
# Configuracion
# --------------------------------------------------------------------------- #

DOCS_POR_DEFECTO = [
    "Bases_2378-57-L126.pdf",
    "decreto_exento_3892_O_243_obligacion_prs.pdf",
    "decreto_exento_3892_Torniquetes.pdf",
]

OUTPUT_ROOT = "Output"

LOCAL_LLM_BASE_URL = os.getenv("LOCAL_LLM_BASE_URL", "http://127.0.0.1:11434/v1")
LOCAL_LLM_API_KEY = os.getenv("LOCAL_LLM_API_KEY", "local")
LOCAL_LLM_MODEL = os.getenv("LOCAL_LLM_MODEL", "qwen3:14b")
LOCAL_LLM_TIMEOUT = float(os.getenv("LOCAL_LLM_TIMEOUT", "180"))
LOCAL_LLM_MAX_CONCURRENCY = max(1, int(os.getenv("LOCAL_LLM_MAX_CONCURRENCY", "2")))

RAG_EMBED_MODEL = os.getenv("RAG_EMBED_MODEL", "BAAI/bge-m3")
RAG_TOP_K = max(2, int(os.getenv("RAG_TOP_K", "8")))
RAG_CHUNK_SIZE = max(600, int(os.getenv("RAG_CHUNK_SIZE", "1800")))
RAG_CHUNK_OVERLAP = max(100, int(os.getenv("RAG_CHUNK_OVERLAP", "220")))
RAG_EMBED_BATCH = max(1, int(os.getenv("RAG_EMBED_BATCH", "16")))
RAG_HYBRID_WEIGHT = min(1.0, max(0.0, float(os.getenv("RAG_HYBRID_WEIGHT", "0.75"))))

SEED_QUERIES = [
    "requisitos de hardware equipos cantidades capacidades especificaciones normas",
    "requisitos de software licencias compatibilidad integracion sistemas operativos bases de datos",
    "instalacion puesta en marcha soporte garantia mantenimiento capacitacion servicios",
    "condiciones tecnicas obligaciones del oferente hitos plazos pruebas aceptacion",
]

_PATRON_ID_LICITACION = re.compile(r"\b\d{3,8}-\d{1,4}-[A-Z]{1,3}\d{1,4}\b")
_MAX_SECCIONES = 60
_MAX_CTX = 4500


# --------------------------------------------------------------------------- #
# Identificacion y normalizacion
# --------------------------------------------------------------------------- #

def _noop(*_args, **_kwargs):
    return None


def _sanitizar(nombre: str) -> str:
    nombre = unicodedata.normalize("NFKD", nombre)
    nombre = nombre.encode("ascii", "ignore").decode("ascii")
    nombre = re.sub(r"[^\w\s.-]", "", nombre).strip()
    nombre = re.sub(r"\s+", "_", nombre)
    return nombre or "licitacion"


def _extraer_texto_pdf(ruta: Path, max_paginas: int = 6) -> str:
    try:
        from pypdf import PdfReader

        lector = PdfReader(str(ruta))
        partes: list[str] = []
        for pagina in lector.pages[:max_paginas]:
            partes.append(pagina.extract_text() or "")
        return "\n".join(partes)
    except Exception:
        return ""


def extraer_id_licitacion(documentos: Iterable[Path]) -> str | None:
    for ruta in documentos:
        ruta = Path(ruta)
        if ruta.suffix.lower() != ".pdf":
            continue
        m = _PATRON_ID_LICITACION.search(ruta.name)
        if m:
            return m.group(0)
        m = _PATRON_ID_LICITACION.search(_extraer_texto_pdf(ruta))
        if m:
            return m.group(0)
    return None


def nombre_licitacion(documentos: Iterable[Path]) -> str:
    lic_id = extraer_id_licitacion(documentos)
    if lic_id:
        return _sanitizar(lic_id)
    documentos = list(documentos)
    if documentos:
        return _sanitizar(Path(documentos[0]).stem)
    return "licitacion"


def _truncar(texto: str, limite: int = _MAX_CTX) -> str:
    if texto and len(texto) > limite:
        return texto[:limite] + "\n[...]"
    return texto


def _limpiar_texto(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    texto = re.sub(r"[ \t]+\n", "\n", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _tokens_busqueda(texto: str) -> set[str]:
    return {
        t
        for t in re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9][\wÁÉÍÓÚÜÑáéíóúüñ.-]{2,}", texto.lower())
        if len(t) >= 3
    }


def _a_markdown(ruta: Path) -> Path:
    dest = ruta.with_name(ruta.stem + ".nlm.txt")
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    try:
        from markitdown import MarkItDown

        resultado = MarkItDown().convert(str(ruta))
        texto = getattr(resultado, "text_content", "") or ""
        if texto.strip():
            dest.write_text(texto, encoding="utf-8")
            return dest
    except Exception:
        pass
    return ruta


def _leer_con_docling(ruta: Path) -> list[tuple[str, str]]:
    try:
        from docling.document_converter import DocumentConverter

        result = DocumentConverter().convert(str(ruta))
        document = getattr(result, "document", None)
        if document is None:
            return []
        if hasattr(document, "export_to_markdown"):
            texto = document.export_to_markdown()
        elif hasattr(document, "export_to_text"):
            texto = document.export_to_text()
        else:
            texto = str(document)
        texto = _limpiar_texto(texto or "")
        return [("documento completo OCR/Docling", texto)] if texto else []
    except Exception:
        return []


def _leer_texto_documento(ruta: Path) -> list[tuple[str, str]]:
    ruta = Path(ruta)
    docling_partes = _leer_con_docling(ruta)
    if docling_partes:
        return docling_partes

    if ruta.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader

            lector = PdfReader(str(ruta))
            partes: list[tuple[str, str]] = []
            for i, pagina in enumerate(lector.pages, start=1):
                texto = _limpiar_texto(pagina.extract_text() or "")
                if texto:
                    partes.append((f"pagina {i}", texto))
            if partes:
                return partes
        except Exception:
            pass

    convertido = _a_markdown(ruta)
    try:
        texto = convertido.read_text(encoding="utf-8")
    except Exception:
        texto = ruta.read_text(encoding="utf-8", errors="ignore")
    texto = _limpiar_texto(texto)
    return [("documento completo", texto)] if texto else []


def _trocear_texto(texto: str, max_chars: int = RAG_CHUNK_SIZE,
                   overlap: int = RAG_CHUNK_OVERLAP) -> list[str]:
    texto = _limpiar_texto(texto)
    if len(texto) <= max_chars:
        return [texto] if texto else []

    bloques = re.split(r"\n\s*\n", texto)
    chunks: list[str] = []
    actual = ""

    def _flush(buffer: str) -> str:
        if not buffer:
            return ""
        if len(buffer) <= max_chars:
            chunks.append(buffer.strip())
            return buffer[-overlap:] if overlap < len(buffer) else buffer
        # Corte por longitud cuando un bloque individual es demasiado grande.
        inicio = 0
        while inicio < len(buffer):
            fin = min(len(buffer), inicio + max_chars)
            chunks.append(buffer[inicio:fin].strip())
            if fin >= len(buffer):
                break
            inicio = max(inicio + 1, fin - overlap)
        return buffer[-overlap:] if overlap < len(buffer) else buffer

    for bloque in bloques:
        bloque = bloque.strip()
        if not bloque:
            continue
        candidato = bloque if not actual else actual + "\n\n" + bloque
        if len(candidato) <= max_chars:
            actual = candidato
            continue
        actual = _flush(actual)
        if len(bloque) > max_chars:
            actual = _flush(bloque)
        else:
            actual = bloque

    if actual:
        _flush(actual)

    # Descartar duplicados triviales y vacios.
    unique = []
    seen = set()
    for chunk in chunks:
        key = chunk[:200]
        if key in seen:
            continue
        seen.add(key)
        if chunk.strip():
            unique.append(chunk.strip())
    return unique


# --------------------------------------------------------------------------- #
# RAG local
# --------------------------------------------------------------------------- #

@dataclass
class Chunk:
    idx: int
    source: str
    page: str
    text: str


class LocalRAGCorpus:
    def __init__(self, documentos: Iterable[Path], embed_model_name: str = RAG_EMBED_MODEL):
        self.documentos = [Path(d) for d in documentos]
        self.embed_model_name = embed_model_name
        self.embedder: SentenceTransformer | None = None
        self.chunks: list[Chunk] = []
        self.embeddings: np.ndarray | None = None
        self._chunk_tokens: list[set[str]] = []

    def _load_embedder(self) -> SentenceTransformer:
        if self.embedder is None:
            self.embedder = SentenceTransformer(self.embed_model_name)
        return self.embedder

    def build(self) -> "LocalRAGCorpus":
        chunks: list[Chunk] = []
        for ruta in self.documentos:
            for page_label, texto in _leer_texto_documento(ruta):
                for frag in _trocear_texto(texto):
                    chunks.append(
                        Chunk(
                            idx=len(chunks) + 1,
                            source=ruta.name,
                            page=page_label,
                            text=frag,
                        )
                    )

        self.chunks = chunks
        self._chunk_tokens = [_tokens_busqueda(c.text) for c in self.chunks]
        if not self.chunks:
            self.embeddings = np.zeros((0, 1), dtype=np.float32)
            return self

        embedder = self._load_embedder()
        texts = [c.text for c in self.chunks]
        vectores = embedder.encode(
            texts,
            batch_size=RAG_EMBED_BATCH,
            normalize_embeddings=True,
            show_progress_bar=False,
        )
        self.embeddings = np.asarray(vectores, dtype=np.float32)
        return self

    def search(self, query: str, top_k: int = RAG_TOP_K) -> list[tuple[Chunk, float]]:
        if self.embeddings is None or not len(self.chunks):
            return []
        embedder = self._load_embedder()
        q = embedder.encode([query], normalize_embeddings=True, show_progress_bar=False)
        qv = np.asarray(q[0], dtype=np.float32)
        dense_scores = self.embeddings @ qv
        query_tokens = _tokens_busqueda(query)
        if query_tokens:
            lexical_scores = np.asarray(
                [
                    len(query_tokens & tokens) / max(1, len(query_tokens))
                    for tokens in self._chunk_tokens
                ],
                dtype=np.float32,
            )
            scores = (RAG_HYBRID_WEIGHT * dense_scores) + ((1 - RAG_HYBRID_WEIGHT) * lexical_scores)
        else:
            scores = dense_scores
        orden = np.argsort(-scores)[: min(top_k, len(self.chunks))]
        return [(self.chunks[i], float(scores[i])) for i in orden]

    def search_many(self, queries: Iterable[str], per_query: int = RAG_TOP_K,
                    max_chunks: int | None = None) -> list[tuple[Chunk, float]]:
        seen: set[int] = set()
        results: list[tuple[Chunk, float]] = []
        for query in queries:
            for chunk, score in self.search(query, per_query):
                if chunk.idx in seen:
                    continue
                seen.add(chunk.idx)
                results.append((chunk, score))
                if max_chunks and len(results) >= max_chunks:
                    return results
        return results


class LocalLLM:
    def __init__(self, base_url: str = LOCAL_LLM_BASE_URL, api_key: str = LOCAL_LLM_API_KEY,
                 model: str = LOCAL_LLM_MODEL, timeout: float = LOCAL_LLM_TIMEOUT):
        self.model = model
        self.client = AsyncOpenAI(base_url=base_url, api_key=api_key, timeout=timeout)

    async def complete(self, system_prompt: str, user_prompt: str,
                       temperature: float = 0.2, max_tokens: int = 1600) -> str:
        resp = await self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=temperature,
            max_tokens=max_tokens,
        )
        contenido = resp.choices[0].message.content if resp.choices else ""
        return (contenido or "").strip()


def _formatear_contexto(chunks: list[tuple[Chunk, float]]) -> str:
    partes: list[str] = []
    for i, (chunk, score) in enumerate(chunks, start=1):
        partes.append(f"[{i}] Fuente: {chunk.source} | {chunk.page} | score={score:.3f}")
        partes.append(chunk.text.strip())
        partes.append("")
    return "\n".join(partes).strip()


def _extraer_refs(texto: str) -> list[int]:
    refs = []
    for n in re.findall(r"\[(\d+)\]", texto):
        v = int(n)
        if v not in refs:
            refs.append(v)
    return refs


def _asegurar_citas_visibles(texto: str, chunks: list[tuple[Chunk, float]]) -> tuple[str, list[int]]:
    refs = [r for r in _extraer_refs(texto) if 1 <= r <= len(chunks)]
    if refs or not chunks:
        return texto.strip(), refs

    lineas: list[str] = []
    for linea in texto.splitlines():
        stripped = linea.strip()
        if (
            not stripped
            or stripped.startswith("#")
            or stripped.startswith("|")
            or re.fullmatch(r"[-:| ]+", stripped)
            or re.search(r"\[\d+\]", stripped)
        ):
            lineas.append(linea)
            continue
        lineas.append(f"{linea} [1]")
    normalizado = "\n".join(lineas).strip()
    return normalizado, ([1] if normalizado else [])


def _citas_desde_refs(chunks: list[tuple[Chunk, float]], refs: list[int]) -> list[dict]:
    citas: list[dict] = []
    for local in refs:
        idx = local - 1
        if idx < 0 or idx >= len(chunks):
            continue
        chunk, _score = chunks[idx]
        citas.append(
            {
                "local": local,
                "texto": chunk.text.strip(),
                "fuente": f"{chunk.source} | {chunk.page}",
            }
        )
    return citas


def _dedupe_items(items: list[dict]) -> list[dict]:
    vistos = set()
    salida: list[dict] = []
    for item in items:
        clave = (
            item["categoria"].lower().strip(),
            re.sub(r"\s+", " ", item["texto"].lower().strip()),
            re.sub(r"\s+", " ", item.get("detalle", "").lower().strip()),
        )
        if clave in vistos:
            continue
        vistos.add(clave)
        salida.append(item)
    return salida


def _parsear_items(texto: str) -> list[dict]:
    items: list[dict] = []
    for linea in texto.splitlines():
        linea = linea.strip()
        if not linea.startswith(("- ", "* ")):
            continue
        m = re.match(r"^[-*]\s*\[(Hardware|Software|Otros)\]\s*(.+)$", linea)
        if not m:
            continue
        categoria = m.group(1)
        resto = m.group(2).strip()
        ref_match = re.findall(r"\[(\d+)\]", resto)
        refs = [int(x) for x in ref_match]
        partes = [p.strip() for p in resto.split("::")]
        texto_item = partes[0] if partes else resto
        detalle = ""
        if len(partes) >= 2:
            detalle = partes[1]
        if len(partes) >= 3 and not refs:
            refs = [int(x) for x in re.findall(r"\[(\d+)\]", partes[2])]
        texto_item = re.sub(r"\[\d+\]", "", texto_item).strip(" -")
        detalle = re.sub(r"\[\d+\]", "", detalle).strip(" -")
        if not texto_item:
            continue
        items.append(
            {
                "categoria": categoria,
                "texto": texto_item,
                "detalle": detalle,
                "refs": refs,
            }
        )
    return items


def _formatear_lista_items(items: list[dict]) -> str:
    grupos = {"Hardware": [], "Software": [], "Otros": []}
    for item in items:
        grupos.setdefault(item["categoria"], []).append(item)

    lineas: list[str] = []
    for categoria in ("Hardware", "Software", "Otros"):
        lineas.append(f"## {categoria}")
        if not grupos.get(categoria):
            lineas.append("- Sin hallazgos claros en esta categoría")
            lineas.append("")
            continue
        for item in grupos[categoria]:
            texto = item["texto"]
            if item.get("detalle"):
                texto = f"{texto} - {item['detalle']}"
            lineas.append(f"- {texto}".rstrip())
        lineas.append("")
    return "\n".join(lineas).strip() + "\n"


def _parsear_indice(texto: str) -> list[str]:
    secciones: list[str] = []
    for linea in texto.splitlines():
        l = linea.strip()
        if not l:
            continue
        m = re.match(r"^\s*(?:\d+[.)]|[-*•]|#+)\s+(.*)$", l)
        titulo = (m.group(1) if m else l).strip()
        titulo = titulo.strip("*").strip()
        if len(titulo) < 3 or len(titulo) > 160:
            continue
        if titulo.lower().startswith((
            "a continuacion",
            "estas son",
            "el indice",
            "aqui",
            "la propuesta",
        )):
            continue
        secciones.append(titulo)
        if len(secciones) >= _MAX_SECCIONES:
            break
    return secciones


async def _ask_local(llm: LocalLLM, corpus: LocalRAGCorpus, query: str, prompt: str,
                     top_k: int = RAG_TOP_K, max_tokens: int = 1600,
                     temperature: float = 0.2) -> tuple[str, list[dict]]:
    chunks = corpus.search(query, top_k=top_k)
    contexto = _formatear_contexto(chunks)
    user_prompt = (
        f"{prompt}\n\n"
        "FUENTES RECUPERADAS:\n"
        f"{contexto}\n\n"
        "Reglas:\n"
        "- Usa solo hechos que aparezcan en las fuentes recuperadas.\n"
        "- Si mencionas un dato, cita con [n].\n"
        "- No inventes ni rellenes vacios.\n"
    )
    texto = await llm.complete(_SYSTEM_PROMPT, user_prompt, temperature=temperature,
                               max_tokens=max_tokens)
    texto, refs = _asegurar_citas_visibles(texto, chunks)
    citas = _citas_desde_refs(chunks, refs)
    return texto.strip(), citas


_SYSTEM_PROMPT = (
    "Eres un analista senior de licitaciones. Tu trabajo es leer documentos "
    "técnicos, extraer requisitos con rigor, proponer respuestas viables y "
    "citar siempre la fuente recuperada. Responde en español, de forma clara "
    "y sin inventar información."
)

_PROMPT_ITEMS = (
    "Extrae items y requerimientos atómicos de esta evidencia para una "
    "licitacion. Devuelve SOLO lineas Markdown en uno de estos formatos:\n"
    "- [Hardware] texto del item\n"
    "- [Software] texto del item\n"
    "- [Otros] texto del item\n"
    "Reglas: un item por linea, sin prosa, sin duplicados obvios, sin "
    "agregar citas en la salida."
)

_PROMPT_INDICE = (
    "A partir de la lista consolidada y las evidencias, propone el indice de "
    "secciones o subsistemas que debe cubrir la propuesta tecnica. Devuelve "
    "solo una lista numerada, una seccion por linea, sin explicaciones."
)


def _prompt_seccion_completa(seccion: str, lista_items: str) -> str:
    return (
        f"En una sola respuesta para la seccion '{seccion}' devuelve exactamente "
        "dos bloques Markdown en este orden:\n"
        "## Requisitos\n"
        "Lista de forma exhaustiva y detallada TODOS los items, equipos, "
        "cantidades, marcas o modelos (si la fuente los indica), normas, "
        "valores numericos, caracteristicas tecnicas y condiciones que los "
        "documentos establecen para esta seccion. Un punto por requisito. No "
        "resumas, no agrupes, no omitas ninguno y no inventes nada. Cita la "
        "fuente cuando corresponda.\n\n"
        "## Propuesta\n"
        "Propuesta tecnica del oferente para esta seccion. Responde CADA "
        "requisito listado. Hardware/equipos: usa tablas Markdown cuando sea "
        "pertinente. Software/servicios/procedimientos: prosa propositiva. "
        "Usa citas [n] en cada afirmacion importante.\n\n"
        "Lista consolidada de items ya extraidos:\n"
        f"{_truncar(lista_items, 3000)}"
    )


async def _extraer_lista_items(llm: LocalLLM, corpus: LocalRAGCorpus) -> str:
    acumulado: list[dict] = []
    sem = asyncio.Semaphore(LOCAL_LLM_MAX_CONCURRENCY)

    async def _run(query: str) -> list[dict]:
        async with sem:
            chunks = corpus.search(query, top_k=max(RAG_TOP_K, 10))
            if not chunks:
                return []
            contexto = _formatear_contexto(chunks)
            user_prompt = (
                f"{_PROMPT_ITEMS}\n\n"
                "EVIDENCIA:\n"
                f"{contexto}\n"
            )
            texto = await llm.complete(_SYSTEM_PROMPT, user_prompt, temperature=0.15,
                                       max_tokens=1400)
            return _parsear_items(texto)

    tareas = [_run(query) for query in SEED_QUERIES]
    resultados = await asyncio.gather(*tareas, return_exceptions=True)
    for res in resultados:
        if isinstance(res, Exception):
            continue
        acumulado.extend(res)
    dedup = _dedupe_items(acumulado)
    return _formatear_lista_items(dedup)


async def _extraer_indice_secciones(llm: LocalLLM, corpus: LocalRAGCorpus,
                                    lista_items: str) -> list[str]:
    query = (
        "indice de secciones de la propuesta tecnica "
        + " ".join(line.strip() for line in lista_items.splitlines()[:80])
    )
    chunks = corpus.search(query, top_k=max(RAG_TOP_K, 10))
    contexto = _formatear_contexto(chunks)
    user_prompt = (
        f"{_PROMPT_INDICE}\n\n"
        "LISTA CONSOLIDADA:\n"
        f"{_truncar(lista_items, 3500)}\n\n"
        "EVIDENCIA:\n"
        f"{contexto}\n"
    )
    texto = await llm.complete(_SYSTEM_PROMPT, user_prompt, temperature=0.15,
                               max_tokens=900)
    secciones = _parsear_indice(texto)
    return secciones or ["Propuesta tecnica general"]


async def _procesar_seccion(llm: LocalLLM, corpus: LocalRAGCorpus, idx: int, seccion: str,
                            ckpt: Path, sem: asyncio.Semaphore,
                            on_progress, n: int, lista_items: str):
    async with sem:
        ruta_req = _ruta_seccion(ckpt, idx, seccion, "01_requisitos")
        ruta_prop = _ruta_seccion(ckpt, idx, seccion, "02_propuesta")

        requisitos = _leer_cuerpo_checkpoint(ruta_req)
        propuesta, citas = _leer_con_citas(ruta_prop)
        if requisitos is not None and propuesta is not None:
            return {
                "seccion": seccion,
                "idx": idx,
                "requisitos": requisitos,
                "propuesta": propuesta,
                "citas": citas,
            }

        on_progress(40 + int(45 * idx / max(1, n)), f"Seccion completa: {seccion}")
        query = f"{seccion} requisitos propuesta tecnica licitacion"
        chunks = corpus.search(query, top_k=max(RAG_TOP_K, 10))
        contexto = _formatear_contexto(chunks)
        user_prompt = (
            f"{_prompt_seccion_completa(seccion, lista_items)}\n\n"
            "EVIDENCIA RELEVANTE:\n"
            f"{contexto}\n"
        )
        texto = await llm.complete(_SYSTEM_PROMPT, user_prompt, temperature=0.2,
                                   max_tokens=1800)
        requisitos, propuesta = _separar_requisitos_y_propuesta(texto)
        if not requisitos and not propuesta:
            propuesta = texto.strip()
        texto, refs = _asegurar_citas_visibles(texto, chunks)
        requisitos, propuesta = _separar_requisitos_y_propuesta(texto)
        if not requisitos and not propuesta:
            propuesta = texto.strip()
        citas = _citas_desde_refs(chunks, refs)

        ruta_req.write_text(f"# {seccion}\n\n{requisitos or ''}", encoding="utf-8")
        ruta_prop.write_text(f"# {seccion}\n\n{propuesta or ''}", encoding="utf-8")
        ruta_prop.with_suffix(".json").write_text(
            json.dumps(citas, ensure_ascii=False), encoding="utf-8"
        )

        return {
            "seccion": seccion,
            "idx": idx,
            "requisitos": requisitos,
            "propuesta": propuesta,
            "citas": citas,
        }


def _separar_requisitos_y_propuesta(texto: str) -> tuple[str, str]:
    if not texto:
        return "", ""

    lineas = texto.splitlines()
    bloque_actual: str | None = None
    buffers = {"requisitos": [], "propuesta": []}

    def _normalizar_encabezado(linea: str) -> str | None:
        m = re.match(
            r"^\s*(?:#{1,6}|\*\*)\s*(requisitos|propuesta)\s*(?:\*\*)?\s*:?\s*$",
            linea.strip(),
            flags=re.IGNORECASE,
        )
        return m.group(1).lower() if m else None

    for linea in lineas:
        encabezado = _normalizar_encabezado(linea)
        if encabezado:
            bloque_actual = encabezado
            continue
        if bloque_actual in buffers:
            buffers[bloque_actual].append(linea)

    requisitos = "\n".join(buffers["requisitos"]).strip()
    propuesta = "\n".join(buffers["propuesta"]).strip()
    if not requisitos and not propuesta:
        partes = re.split(r"(?im)^\s*#{1,6}\s*propuesta\s*:?\s*$", texto, maxsplit=1)
        if len(partes) == 2:
            cabeza, cola = partes
            req_parts = re.split(r"(?im)^\s*#{1,6}\s*requisitos\s*:?\s*$", cabeza, maxsplit=1)
            if len(req_parts) == 2:
                requisitos = req_parts[1].strip()
            propuesta = cola.strip()
    return requisitos, propuesta


def _ruta_seccion(ckpt: Path, idx: int, seccion: str, sub: str) -> Path:
    slug = f"{idx + 1:02d}_{_sanitizar(seccion)[:40]}"
    return Path(ckpt) / sub / f"{slug}.md"


def _leer_si_existe(ruta: Path) -> str | None:
    p = Path(ruta)
    if p.exists() and p.stat().st_size > 0:
        return p.read_text(encoding="utf-8")
    return None


def _leer_cuerpo_checkpoint(ruta: Path) -> str | None:
    texto = _leer_si_existe(ruta)
    if texto is None:
        return None
    lineas = texto.splitlines()
    if lineas and lineas[0].startswith("# "):
        return "\n".join(lineas[2:] if len(lineas) > 1 and lineas[1] == "" else lineas[1:])
    return texto


def _leer_con_citas(ruta_md: Path) -> tuple[str | None, list[dict]]:
    p = Path(ruta_md)
    if not (p.exists() and p.stat().st_size > 0):
        return None, []
    texto = p.read_text(encoding="utf-8")
    lineas = texto.splitlines()
    if lineas and lineas[0].startswith("# "):
        texto = "\n".join(lineas[2:] if len(lineas) > 1 and lineas[1] == "" else lineas[1:])
    if not texto.strip():
        return None, []
    cj = p.with_suffix(".json")
    citas = json.loads(cj.read_text(encoding="utf-8")) if cj.exists() else []
    return texto, citas


def _renumerar_citas_global(datos: list[dict]):
    global_n = 1
    for d in datos:
        for c in d.get("citas", []):
            c["global"] = global_n
            global_n += 1


def _parsear_tabla_md(texto: str) -> list[list[str]]:
    filas = []
    for linea in texto.splitlines():
        l = linea.strip()
        if not l.startswith("|"):
            continue
        celdas = [c.strip() for c in l.strip("|").split("|")]
        if celdas and all(set(c) <= set("-: ") and c != "" for c in celdas):
            continue
        filas.append(celdas)
    return filas


def _sombrear_celda(celda, color_hex: str = "2E5EA8"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = celda._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _agregar_tabla(doc, filas):
    from docx.shared import Inches, Pt, RGBColor

    ncols = max(len(f) for f in filas)
    tabla = doc.add_table(rows=0, cols=ncols)
    tabla.style = "Table Grid"

    if ncols == 2:
        anchos = [Inches(2.2), Inches(4.0)]
    elif ncols == 3:
        anchos = [Inches(2.0), Inches(3.2), Inches(1.0)]
    else:
        anchos = [Inches(6.2 / ncols)] * ncols

    for i, fila in enumerate(filas):
        row = tabla.add_row()
        for j in range(ncols):
            valor = fila[j].replace("**", "") if j < len(fila) else ""
            celda = row.cells[j]
            celda.width = anchos[j] if j < len(anchos) else Inches(1.5)
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(valor)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _sombrear_celda(celda)

    doc.add_paragraph()


def _agregar_linea(doc, linea: str):
    from docx.shared import Pt

    texto = linea.rstrip()
    if not texto.strip():
        return

    stripped = texto.lstrip()

    if stripped.startswith("#"):
        nivel = len(stripped) - len(stripped.lstrip("#"))
        titulo = stripped[nivel:].strip()
        doc.add_heading(titulo, level=min(nivel, 4))
        return

    if stripped[:2] in ("- ", "* ") or stripped.startswith("• "):
        p = doc.add_paragraph(style="List Bullet")
        _runs_con_negrita(p, stripped[2:].strip())
        return

    if _es_lista_numerada(stripped):
        contenido = stripped.split(None, 1)[1] if " " in stripped else stripped
        p = doc.add_paragraph(style="List Number")
        _runs_con_negrita(p, contenido)
        return

    p = doc.add_paragraph()
    _runs_con_negrita(p, stripped)


def _es_lista_numerada(texto: str) -> bool:
    partes = texto.split()
    cabeza = partes[0] if partes else ""
    return bool(cabeza) and cabeza[:-1].isdigit() and cabeza[-1:] in (".", ")")


def _runs_con_negrita(parrafo, texto: str):
    partes = texto.split("**")
    for i, parte in enumerate(partes):
        if not parte:
            continue
        run = parrafo.add_run(parte)
        if i % 2 == 1:
            run.bold = True


def _renderizar_contenido(doc, texto: str):
    en_tabla = False
    bloque_tabla: list[str] = []

    for linea in texto.splitlines():
        es_fila_tabla = linea.strip().startswith("|")
        if es_fila_tabla:
            if not en_tabla:
                en_tabla = True
                bloque_tabla = []
            bloque_tabla.append(linea)
            continue

        if en_tabla:
            filas = _parsear_tabla_md("\n".join(bloque_tabla))
            if filas:
                _agregar_tabla(doc, filas)
            en_tabla = False
            bloque_tabla = []

        _agregar_linea(doc, linea)

    if en_tabla and bloque_tabla:
        filas = _parsear_tabla_md("\n".join(bloque_tabla))
        if filas:
            _agregar_tabla(doc, filas)


def _nuevo_doc(titulo: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(titulo)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fecha = _dt.date.today().strftime("%d-%m-%Y")
    srun = sub.add_run(f"Propuesta del oferente · {fecha}")
    srun.italic = True
    srun.font.size = Pt(11)
    srun.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    doc.add_paragraph()
    return doc


def guardar_docx(titulo: str, contenido: str, ruta_salida: str):
    doc = _nuevo_doc(titulo)
    _renderizar_contenido(doc, contenido)
    doc.save(ruta_salida)
    print(f"-> Documento guardado en: {ruta_salida}")


def generar_docx_desde_resultado(result_json: dict, carpeta: Path) -> Path:
    lic = result_json["licitacion"]
    sin_dato = "_Esta seccion no se pudo generar._"

    cuerpo = ["## Resumen ejecutivo", result_json.get("resumen") or sin_dato, ""]
    for sec in result_json.get("secciones", []):
        cuerpo.append(f"## {sec['titulo']}")
        cuerpo.append(sec.get("propuesta") or sin_dato)
        cuerpo.append("")

    ruta = carpeta / f"Propuesta_Tecnica_{lic}.docx"
    guardar_docx("Propuesta Tecnica del Oferente", "\n".join(cuerpo), str(ruta))
    return ruta


async def _construir_corpus(documentos: Iterable[Path]) -> LocalRAGCorpus:
    corpus = LocalRAGCorpus(documentos)
    return await asyncio.to_thread(corpus.build)


async def _construir_llm() -> LocalLLM:
    try:
        return LocalLLM()
    except Exception as exc:
        raise RuntimeError(
            "No se pudo inicializar el cliente LLM local. "
            "Verifica LOCAL_LLM_BASE_URL, LOCAL_LLM_MODEL y que tu servidor "
            "OpenAI-compatible este levantado."
        ) from exc


async def _ejecutar_pipeline(documentos, nombre_notebook, carpeta, on_progress,
                             storage_path=None):
    del nombre_notebook, storage_path

    ckpt = Path(carpeta) / "checkpoints"
    (ckpt / "01_requisitos").mkdir(parents=True, exist_ok=True)
    (ckpt / "02_propuesta").mkdir(parents=True, exist_ok=True)

    lista_items = _leer_si_existe(ckpt / "00_lista_items.md")
    indice_raw = _leer_si_existe(ckpt / "00_indice_secciones.md")
    resumen = _leer_si_existe(ckpt / "04_resumen.md")
    secciones = _parsear_indice(indice_raw) if indice_raw else None

    necesita_llm = lista_items is None or indice_raw is None or resumen is None
    if secciones:
        for i, s in enumerate(secciones):
            if (
                _leer_cuerpo_checkpoint(_ruta_seccion(ckpt, i, s, "01_requisitos")) is None
                or _leer_con_citas(_ruta_seccion(ckpt, i, s, "02_propuesta"))[0] is None
            ):
                necesita_llm = True
                break
    else:
        necesita_llm = True

    documentos = [Path(d) for d in documentos]
    if not necesita_llm:
        datos = []
        for i, s in enumerate(secciones or []):
            prop, citas = _leer_con_citas(_ruta_seccion(ckpt, i, s, "02_propuesta"))
            datos.append(
                {
                    "seccion": s,
                    "idx": i,
                    "requisitos": _leer_cuerpo_checkpoint(_ruta_seccion(ckpt, i, s, "01_requisitos")),
                    "propuesta": prop,
                    "citas": citas,
                }
            )
        _renumerar_citas_global(datos)
        return {"lista_items": lista_items or "", "resumen": resumen or "", "secciones": datos}

    on_progress(3, "Construyendo corpus local")
    corpus = await _construir_corpus(documentos)
    llm = await _construir_llm()

    on_progress(12, "Extrayendo lista de items")
    if lista_items is None:
        lista_items = await _extraer_lista_items(llm, corpus)
        (ckpt / "00_lista_items.md").write_text(lista_items, encoding="utf-8")

    on_progress(25, "Proponiendo indice de secciones")
    if indice_raw is None:
        secciones = await _extraer_indice_secciones(llm, corpus, lista_items)
        indice_raw = "\n".join(f"{i + 1}. {s}" for i, s in enumerate(secciones))
        (ckpt / "00_indice_secciones.md").write_text(indice_raw, encoding="utf-8")
    else:
        secciones = _parsear_indice(indice_raw)

    if not secciones:
        secciones = ["Propuesta tecnica general"]

    sem = asyncio.Semaphore(LOCAL_LLM_MAX_CONCURRENCY)
    tareas = [
        _procesar_seccion(llm, corpus, idx, seccion, ckpt, sem, on_progress,
                          len(secciones), lista_items)
        for idx, seccion in enumerate(secciones)
    ]
    resultados_raw = await asyncio.gather(*tareas, return_exceptions=True)

    datos = []
    for i, res in enumerate(resultados_raw):
        if isinstance(res, Exception):
            datos.append(
                {
                    "seccion": secciones[i],
                    "idx": i,
                    "requisitos": None,
                    "propuesta": None,
                    "citas": [],
                }
            )
        else:
            datos.append(res)

    _renumerar_citas_global(datos)

    on_progress(90, "Redactando resumen ejecutivo")
    if resumen is None:
        query = "resumen ejecutivo propuesta tecnica licitacion " + " ".join(secciones[:10])
        chunks = corpus.search(query, top_k=max(RAG_TOP_K, 10))
        contexto = _formatear_contexto(chunks)
        user_prompt = (
            "Redacta un resumen ejecutivo breve de maximo 200 palabras de la "
            "propuesta tecnica del oferente. Destaca el alcance general y los "
            "principales sistemas o equipos ofertados. Cita con [n] cuando sea "
            "relevante.\n\n"
            "EVIDENCIA:\n"
            f"{contexto}\n"
        )
        resumen = await llm.complete(_SYSTEM_PROMPT, user_prompt, temperature=0.2,
                                     max_tokens=350)
        (ckpt / "04_resumen.md").write_text(resumen, encoding="utf-8")

    return {"lista_items": lista_items, "resumen": resumen, "secciones": datos}


async def generar_propuesta(documentos, output_root=OUTPUT_ROOT, on_progress=_noop,
                            nombre_notebook=None, storage_path=None):
    documentos = [Path(d) for d in documentos]
    for d in documentos:
        if not d.exists():
            raise FileNotFoundError(f"No se encontro el documento: {d}")

    on_progress(2, "Identificando la licitacion")
    lic = nombre_licitacion(documentos)
    if nombre_notebook is None:
        nombre_notebook = f"Licitacion {lic}"

    carpeta = Path(output_root) / lic
    carpeta.mkdir(parents=True, exist_ok=True)

    res = await _ejecutar_pipeline(
        documentos,
        nombre_notebook,
        carpeta,
        on_progress,
        storage_path=storage_path,
    )

    on_progress(95, "Generando documentos")
    fecha = _dt.date.today().strftime("%d/%m/%Y")
    sin_dato = (
        "_Esta seccion no se pudo generar en esta ejecucion. Re-ejecutar el "
        "proceso para reintentarla; el resto se conserva desde los checkpoints._"
    )

    result_json = {
        "version": 1,
        "licitacion": lic,
        "fecha": _dt.datetime.now().isoformat(),
        "lista_items": res["lista_items"] or "",
        "resumen": res["resumen"] or "",
        "secciones": [
            {
                "idx": d["idx"],
                "titulo": d["seccion"],
                "propuesta": d["propuesta"] or "",
                "citas": d.get("citas", []),
            }
            for d in res["secciones"]
        ],
    }
    ruta_json = carpeta / "result.json"
    ruta_json.write_text(json.dumps(result_json, ensure_ascii=False, indent=2),
                         encoding="utf-8")

    cuerpo_propuesta = ["## Resumen ejecutivo", res["resumen"] or sin_dato, ""]
    for d in res["secciones"]:
        cuerpo_propuesta.append(f"## {d['seccion']}")
        cuerpo_propuesta.append(d["propuesta"] or sin_dato)
        cuerpo_propuesta.append("")

    ruta_propuesta = carpeta / f"Propuesta_Tecnica_{lic}.docx"
    guardar_docx("Propuesta Tecnica del Oferente", "\n".join(cuerpo_propuesta),
                 str(ruta_propuesta))

    ruta_items = carpeta / f"Lista_Items_{lic}.docx"
    guardar_docx("Items Requeridos por la Licitacion", res["lista_items"] or sin_dato,
                 str(ruta_items))

    on_progress(100, "Proceso finalizado")
    return {
        "licitacion": lic,
        "carpeta": str(carpeta),
        "result_path": str(ruta_json),
        "archivos": [
            {
                "id": "propuesta",
                "title": "Propuesta tecnica",
                "description": "Propuesta del oferente, seccion por seccion y punto por punto.",
                "filename": ruta_propuesta.name,
                "path": str(ruta_propuesta),
                "date": fecha,
            },
            {
                "id": "items",
                "title": "Lista de items",
                "description": "Items requeridos, agrupados en Hardware, Software y Otros.",
                "filename": ruta_items.name,
                "path": str(ruta_items),
                "date": fecha,
            },
        ],
    }


async def reprocessar_secciones(carpeta, section_indices, documentos, on_progress=_noop,
                                storage_path=None):
    del storage_path
    carpeta = Path(carpeta)
    ruta_json = carpeta / "result.json"
    resultado_actual = json.loads(ruta_json.read_text(encoding="utf-8"))

    ckpt = carpeta / "checkpoints"
    secciones_info = resultado_actual["secciones"]
    documentos = [Path(d) for d in documentos]

    on_progress(5, "Reconstruyendo corpus local")
    corpus = await _construir_corpus(documentos)
    llm = await _construir_llm()
    lista_items = _leer_si_existe(ckpt / "00_lista_items.md") or resultado_actual.get("lista_items", "")

    for idx in section_indices:
        if idx >= len(secciones_info):
            continue
        seccion = secciones_info[idx]["titulo"]
        ruta_req = _ruta_seccion(ckpt, idx, seccion, "01_requisitos")
        ruta_prop = _ruta_seccion(ckpt, idx, seccion, "02_propuesta")
        for p in [ruta_req, ruta_prop, ruta_prop.with_suffix(".json")]:
            if p.exists():
                p.unlink()

    sem = asyncio.Semaphore(LOCAL_LLM_MAX_CONCURRENCY)
    n = len(secciones_info)
    indices_validos = [idx for idx in section_indices if idx < n]
    tareas = [
        _procesar_seccion(
            llm,
            corpus,
            idx,
            secciones_info[idx]["titulo"],
            ckpt,
            sem,
            on_progress,
            n,
            lista_items,
        )
        for idx in indices_validos
    ]
    resultados = await asyncio.gather(*tareas, return_exceptions=True)

    for i, idx in enumerate(indices_validos):
        res = resultados[i]
        if isinstance(res, Exception):
            continue
        secciones_info[idx] = {
            "idx": idx,
            "titulo": res["seccion"],
            "propuesta": res["propuesta"] or "",
            "citas": res.get("citas", []),
        }

    _renumerar_citas_global(secciones_info)

    v_anterior = resultado_actual.get("version", 1)
    (carpeta / f"result_v{v_anterior}.json").write_text(
        ruta_json.read_text(encoding="utf-8"), encoding="utf-8"
    )

    resultado_actual["version"] = v_anterior + 1
    resultado_actual["secciones"] = secciones_info
    ruta_json.write_text(json.dumps(resultado_actual, ensure_ascii=False, indent=2),
                         encoding="utf-8")

    return resultado_actual


def main():
    parser = argparse.ArgumentParser(
        description="Procesa una licitacion con RAG local y genera los .docx"
    )
    parser.add_argument("documentos", nargs="*", default=DOCS_POR_DEFECTO,
                        help="Documentos a subir (por defecto: los PDFs del repo).")
    parser.add_argument("--output", default=OUTPUT_ROOT,
                        help=f"Carpeta raiz de salida (por defecto: {OUTPUT_ROOT}).")
    args = parser.parse_args()

    def progreso(p, m):
        print(f"  [{p:3d}%] {m}")

    resultado = asyncio.run(
        generar_propuesta(args.documentos, output_root=args.output, on_progress=progreso)
    )

    print(f"\n-> Licitacion: {resultado['licitacion']}")
    print(f"-> Carpeta:   {resultado['carpeta']}")
    print(f"-> JSON:      {resultado['result_path']}")
    for a in resultado["archivos"]:
        print(f"  - {a['title']}: {a['path']}")


if __name__ == "__main__":
    main()
